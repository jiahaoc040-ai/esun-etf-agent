import copy
import json

import docx
import pytest

from esun_agent import strategy_doc as sd
from esun_agent.agent.observations import BASIS_LABELS
from esun_agent.strategy.declaration import load_declaration

SUMMARY = {
    "first": "2025-04-01", "oos_start": "2026-08-01", "length": 24, "step": 2, "universe_dates": ["2025-01-02", "2026-10-05"],
    "strict": {"n": 171, "excess": {"n": 171, "mean": -0.012, "p5": -0.133, "p95": 0.089, "p_gt_5": 0.19, "p_gt_10": 0.04, "p_lt_m10": 0.12},
        "mdd": {"mean": -0.073}, "turnover": {"mean": 3.2}, "cost": {"mean": 0.018}, "sum_v_holdings": 0, "sum_v_cash": 0,
        "sum_ap_fail_days": 0, "sum_drift_days": 82},
    "strict_oos": {"n": 11, "excess": {"n": 11, "mean": -0.112, "p5": -0.2, "p95": -0.05, "p_gt_5": 0.0, "p_gt_10": 0.0, "p_lt_m10": 0.64},
        "mdd": {"mean": -0.09}, "turnover": {"mean": 3.2}, "cost": {"mean": 0.018}, "sum_v_holdings": 0, "sum_v_cash": 0,
        "sum_ap_fail_days": 0, "sum_drift_days": 82},
    "hysteresis": {"n": 171, "excess": {"n": 171, "mean": 0.004, "p5": -0.12, "p95": 0.13, "p_gt_5": 0.26, "p_gt_10": 0.13, "p_lt_m10": 0.09},
        "mdd": {"mean": -0.06}, "turnover": {"mean": 3.2}, "cost": {"mean": 0.018}, "sum_v_holdings": 0, "sum_v_cash": 0,
        "sum_ap_fail_days": 0, "sum_drift_days": 82},
    "tsmc_fixed": {"n": 171, "excess": {"n": 171, "mean": -0.013, "p5": -0.13, "p95": 0.09, "p_gt_5": 0.13, "p_gt_10": 0.04, "p_lt_m10": 0.12},
        "mdd": {"mean": -0.07}, "turnover": {"mean": 3.2}, "cost": {"mean": 0.018}, "sum_v_holdings": 0, "sum_v_cash": 0,
        "sum_ap_fail_days": 0, "sum_drift_days": 82},
    "base_a": {"n": 171, "excess": {"n": 171, "mean": -0.009, "p5": -0.037, "p95": 0.031, "p_gt_5": 0.02, "p_gt_10": 0.0, "p_lt_m10": 0.0},
        "mdd": {"mean": -0.05}, "turnover": {"mean": 3.2}, "cost": {"mean": 0.018}, "sum_v_holdings": 0, "sum_v_cash": 0,
        "sum_ap_fail_days": 0, "sum_drift_days": 82},
}


def texts(d):
    out = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                out.append(c.text)
    return "\n".join(out)


def build(tmp_path, **kw):
    return docx.Document(str(sd.build_strategy_doc(tmp_path / "doc.docx", summary=SUMMARY, **kw)))


def test_form_cells_come_from_declaration(tmp_path):
    decl = load_declaration()
    d = build(tmp_path)
    cells = [r.cells[1].text for r in d.tables[0].rows]
    assert cells == [decl["etf"]["name"], decl["theme"]["title"], decl["philosophy_statement"]]
    assert cells[0].startswith("主動") and len(cells[0]) <= 20 and len(cells[1]) <= 50 and 100 <= len(cells[2]) <= 300
    # 版型沿用官方範本：標題欄文字與頁首不變、表單欄位使用標楷體
    assert [r.cells[0].text for r in d.tables[0].rows][0].startswith("ETF 名稱")
    assert "參賽隊伍" in "".join(p.text for p in d.sections[0].header.paragraphs)
    run = d.tables[0].rows[0].cells[1].paragraphs[0].runs[0]
    assert run.font.name == "標楷體"


def test_refuses_noncompliant_declaration(tmp_path):
    decl = copy.deepcopy(load_declaration())
    decl["etf"]["name"] = "台灣核心優選 AI 主動 ETF"
    with pytest.raises(sd.DocError, match="開頭"):
        sd.build_strategy_doc(tmp_path / "x.docx", decl=decl, summary=SUMMARY)
    assert not (tmp_path / "x.docx").exists()
    decl = copy.deepcopy(load_declaration())
    decl["philosophy_statement"] = "短"
    with pytest.raises(sd.DocError, match="philosophy_statement"):
        sd.build_strategy_doc(tmp_path / "y.docx", decl=decl, summary=SUMMARY)


def test_form_only(tmp_path):
    d = build(tmp_path, form_only=True)
    assert "補充說明" not in texts(d) and len(d.tables) == 1


def test_supplement_content_and_numbers(tmp_path):
    decl = load_declaration()
    t = texts(build(tmp_path))
    for h in ("補充說明", "一、投組結構", "二、超額來源", "三、換股紀律", "四、代理回測摘要", "五、降級保底", "六、每日 D-Plan"):
        assert h in t
    r = decl["portfolio_rules"]
    # 規則數字來自宣告檔
    for x in (r["w_high"], r["w_mid"], r["cap_tsmc"], r["ap_target"], r["ap_recovery_target"], r["cash_high"]):
        assert f"{x * 100:.0f}%" in t, x
    assert f"{r['n_conv_min']}–{r['n_conv_max']} 檔" in t and f"{r['n_total_min']}–{r['n_total_max']} 檔" in t
    for k in ("min_hold_days", "max_swaps", "stop_loss", "event_blackout"):
        assert decl["discipline"][k] in t
    for b in decl["selection"]["allowed_bases"]:
        assert b["label"] in t
        for a in b["authority"]:
            assert a in t
    assert decl["fallback"]["name"] in t and "凍結" in t
    # 回測數字來自摘要，不是寫死
    for s_ in ("19%", "4%", "12%", "-13.3%", "+8.9%", "2026-08-01", "-11.2%", "24 個交易日"):
        assert s_ in t, s_
    assert "前視與倖存者偏誤" in t and "不證明 Agent" in t                      # 名單偏誤與限制必須寫明
    assert "連續 2 日" in t and "FOMC" in t


def test_every_section_maps_to_dplan_fields(tmp_path):
    d = build(tmp_path)
    boxes = [t.rows[0].cells[0].text for t in d.tables if t.rows[0].cells[0].text.startswith("對應每日 D-Plan")]
    assert len(boxes) >= 4
    for b in boxes:
        assert "market_view" in b or "inferences" in b
    allt = "\n".join(boxes)
    for k in ("market_view", "inferences", "decisions", "observations", "no_trade_decisions"):
        assert k in allt


def test_mapping_matches_agent_code(env=None):
    """文件裡引用的 observation topic／依據 id 必須是 Agent 程式真的會產生的。"""
    decl = load_declaration()
    assert {b["id"] for b in decl["selection"]["allowed_bases"]} == set(BASIS_LABELS)
    assert {t for t, _ in sd.MARKET_TOPICS} == {"tw_market", "inst_flow", "taifex_night", "us_overnight"}
    from esun_agent.agent import observations as ob
    from esun_agent.data.prices import load_panel
    panel = load_panel()
    ev = ob.build_evidence(panel, "2026-10-02", inputs={
        "taifex": {"url": "https://x.example/t", "content_as_of": "2026-10-02T05:00:00+08:00", "close": 1.0, "chg_pct": 0.1},
        "us": {"url": "https://x.example/u", "content_as_of": "2026-10-02T04:00:00+08:00", "quotes": {"SOX": {"chg_pct": 1.0}}}})
    assert {t for t, _ in sd.MARKET_TOPICS} <= {o["topic"] for o in ev.observations}
    import inspect
    from esun_agent.agent import conviction_agent
    src = inspect.getsource(ob) + inspect.getsource(conviction_agent)
    for topic, _ in sd.STRATEGY_TOPICS:
        assert f'"{topic}"' in src, topic


def test_cli(tmp_path, capsys):
    import importlib.util
    from pathlib import Path
    spec = importlib.util.spec_from_file_location("make_doc", Path(__file__).parent.parent / "scripts" / "make_strategy_doc.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    assert mod.main(["--out", str(tmp_path / "o.docx"), "--form-only"]) == 0
    assert (tmp_path / "o.docx").exists()
    missing = sd.SUMMARY_PATH
    if not missing.exists():
        assert mod.main(["--out", str(tmp_path / "p.docx")]) == 1


def test_real_summary_file_builds_when_present(tmp_path):
    if not sd.SUMMARY_PATH.exists():
        pytest.skip("尚未產生 docs/conviction_backtest.json")
    S = sd.load_summary()
    d = docx.Document(str(sd.build_strategy_doc(tmp_path / "real.docx", summary=S)))
    assert pct_in(S["strict"]["excess"]["p5"], texts(d)) and pct_in(S["strict_oos"]["excess"]["mean"], texts(d))


def pct_in(x, t):
    return sd.pct(x) in t
