"""ETF 投資策略說明文件（docx）產生器。

版型沿用主辦方的「ETF 投資策略說明文件_參賽隊伍繳交格式_v1.docx」：第一頁三欄表單（ETF 名稱、投資主題、投資理念）
由 config/strategy_declaration.json 填入（名稱 ≤20 字且開頭須為「主動」、主題 ≤50 字、理念 100～300 字，不符就不產檔）。
表單之後是「補充說明」（form_only=True 時省略）：投組結構、候選池與 Agent 選股、換股紀律、代理回測摘要、降級保底、
與每日 D-Plan 的對應。所有文字內容都來自宣告檔（v2：集中主題＋事件驅動）與 docs/conviction_backtest.json，
不另寫一份；數字隨回測結果更新。
"""
from __future__ import annotations

import json
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from .config import ROOT
from .strategy.declaration import load_declaration, validate_submission

TEMPLATE = ROOT / "docs" / "competition" / "玉山挑戰賽_ETF投資策略說明文件_參賽隊伍繳交格式_v1.docx"
SUMMARY_PATH = ROOT / "docs" / "conviction_backtest.json"
OUTPUT_PATH = ROOT / "output" / "strategy_doc.docx"
FORM_FONT = "標楷體"
BODY_FONT = "微軟正黑體"
BLUE = RGBColor(0x1F, 0x4E, 0x79)

# 宣告 → 每日 D-Plan 的對應（observation.topic 與宣告檔的依據 id 相同；程式在 esun_agent/agent/ 內檢查）
MARKET_TOPICS = [("tw_market", "台股整體（名單內 150 檔等權報酬、漲跌家數、成交值）"), ("inst_flow", "外資、投信買賣超金額"),
                 ("taifex_night", "台指期夜盤"), ("us_overnight", "美股／ADR 收盤")]
STRATEGY_TOPICS = [("candidate_screen", "候選分數與組成（排序用，不可單獨作為依據）"),
                   ("trust_flow", "投信近 5／10 日淨買超占成交值、連續買超天數"),
                   ("position_pnl", "高信心部位的持有天數與相對成本報酬（停損、最短持有）"),
                   ("filler_universe", "市值前 30 大的湊檔數部位"),
                   ("event_calendar", "FOMC／NVIDIA 財報／美國大選排程（事件前一交易日）"),
                   ("weights_vs_target", "現權重與目標權重的偏離（無交易帶）")]


class DocError(ValueError):
    pass


def pct(x: float, d: int = 1, sign: bool = True) -> str:
    return f"{x * 100:+.{d}f}%" if sign else f"{x * 100:.{d}f}%"


def load_summary(path: Path | None = None) -> dict:
    p = path or SUMMARY_PATH
    if not p.exists():
        raise DocError(f"找不到回測摘要 {p}，請先執行 python scripts/run_conviction_backtest.py")
    return json.loads(p.read_text(encoding="utf-8"))


# ----------------------------------------------------------------------------- docx 小工具

def _font(run, name=BODY_FONT, size=None, bold=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rf.set(qn(a), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _shade(cell, fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    _font(p.add_run(text), BODY_FONT, 14 if level == 1 else 11.5, True, BLUE)
    return p


def para(doc, text, *, bold=False, size=10, italic=False, color=None, after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    _font(r, BODY_FONT, size, bold, color)
    r.italic = italic
    return p


def bullets(doc, items, size=10):
    for t in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(14)
        p.paragraph_format.first_line_indent = Pt(-10)
        p.paragraph_format.space_after = Pt(2)
        _font(p.add_run("• " + t), BODY_FONT, size)


def table(doc, header, rows, widths=None, size=9):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        _font(c.paragraphs[0].add_run(h), BODY_FONT, size, True, BLUE)
        _shade(c, "F2F6FA")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            _font(cells[i].paragraphs[0].add_run(str(v)), BODY_FONT, size)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Pt(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def mapping_box(doc, lines: list[str]):
    """每段宣告結尾的「D-Plan 對應」框：說明這段宣告在每日 D-Plan 的哪些欄位出現。"""
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.rows[0].cells[0]
    _shade(c, "F7F9FC")
    c.text = ""
    _font(c.paragraphs[0].add_run("對應每日 D-Plan"), BODY_FONT, 9, True, BLUE)
    for ln in lines:
        p = c.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        _font(p.add_run("• " + ln), BODY_FONT, 9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ----------------------------------------------------------------------------- 表單

def _fill_form(doc, decl: dict) -> None:
    t = doc.tables[0]
    values = [decl["etf"]["name"], decl["theme"]["title"], decl["philosophy_statement"]]
    if len(t.rows) != 3:
        raise DocError("官方範本表單應為 3 列（名稱／主題／理念）")
    for row, v in zip(t.rows, values):
        cell = row.cells[1]
        p = cell.paragraphs[0]
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
        _font(p.add_run(v), FORM_FONT, 10.5)


# ----------------------------------------------------------------------------- 補充說明各節

def _sec_structure(doc, decl, S):
    r = decl["portfolio_rules"]
    heading(doc, "補充說明")
    para(doc, "以下內容補充說明上表的宣告，並標示每段宣告將如何出現在每日 D-Plan（market_view、inferences、decisions）。", size=9, italic=True)
    heading(doc, "一、投組結構：集中的高信心部位＋湊檔數部位", 2)
    para(doc, decl["theme"]["summary"], size=10)
    table(doc, ["部位", "規則"], [
        ["投資範圍", decl["etf"]["market"]],
        ["高信心部位", f"{r['n_conv_min']}–{r['n_conv_max']} 檔，由 AI Agent 從候選池挑選並給信心等級：高 = {pct(r['w_high'], 0, False)}、"
                     f"中 = {pct(r['w_mid'], 1, False)}（2330：高 {pct(r['w_tsmc_high'], 0, False)}、中 {pct(r['w_tsmc_mid'], 0, False)}）；"
                     "名目權重合計過大時等比例縮小、過小時在上限內等比例放大"],
        ["單檔上限", f"{pct(r['cap_other'], 0, False)}（官方 10%，留 2% 給價格漂移：買進當天漲停且淨值下跌仍不超過 10%），2330 為 {pct(r['cap_tsmc'], 0, False)}（官方 25%）；"
                   f"前日收盤超過 {pct(r['guard_other'], 1, False)}（2330 {pct(r['guard_tsmc'], 0, False)}）當日減回"],
        ["湊檔數部位", f"市值前 {r['filler_pool']} 大、且不在高信心名單者，取市值最大的若干檔各 {pct(r['filler_min'], 0, False)}–{pct(r['filler_max'], 0, False)}，"
                     f"使總持股 {r['n_total_min']}–{r['n_total_max']} 檔（官方 20–30；內部至少比 20 檔多 3 檔緩衝）；"
                     f"出現在超過 {r['filler_max_overlap']:.0%} 主動型 ETF 前 15 大的股票排在後面（低重疊優先）"],
        ["現金", f"目標 {pct(r['cash_target'], 0, False)}，區間 {pct(r['cash_low'], 0, False)}–{pct(r['cash_high'], 0, False)}；"
               "委託後預估現金比 ≥ 2% ＋ 0.03 × 買進金額占淨值（成交均價可能高於前日收盤），不足時等比例縮減買進"],
        ["Active Share", f"我方前 15 大 vs 每檔主動型 ETF 前 15 大，官方 ≥ 15%、內部目標 ≥ {pct(r['ap_target'], 0, False)}；"
                         "不足時先把重疊的湊檔股換成不重疊者，再把重疊最大的高信心部位每次降 1.5%（最低 3%）；仍不足則退回 Agent 換股"],
        ["Active Share 硬規則", f"前日收盤 < {pct(r['ap_trigger'], 0, False)} → 當日強制再平衡到 ≥ {pct(r['ap_target'], 0, False)}；"
                                f"前日收盤已 < 15% → 當日補到 ≥ {pct(r['ap_recovery_target'], 0, False)}，避免連續 2 日低於官方門檻"],
        ["交易", f"每日依名單重算目標權重；既有持股差距 < {pct(r['band'], 1, False)}（無交易帶）不交易；單日異動超過 30 檔（schema 上限）時分兩天完成"],
    ], widths=[85, 380])
    mapping_box(doc, [
        "decisions：target_weight 由信心等級機械換算（高 8%／中 5.5%），湊檔股各 1–2%；委託一律由官方公式 derive_orders 推導，Agent 不決定權重或股數。",
        "inferences：每檔高信心部位一則推論（引用 Agent 選的 observation）；湊檔數部位一則推論（引用 filler_universe）；"
        "未偏離無交易帶者以「差距在無交易帶內」推論寫入 no_trade_decisions。",
        "market_view：posture 與現金區間由實際委託推得，不由 Agent 自填；regime／stance 須與淨買賣方向一致。",
    ])


def _sec_selection(doc, decl, S):
    c = decl["candidate_rules"]
    heading(doc, "二、超額來源：候選池＋AI Agent 依公開事件選股", 2)
    para(doc, "程式每日從 150 檔用資料篩出候選池，Agent 只能從候選池（或目前名單）挑選；候選分數負責排序，"
              "入選理由必須是具名的公開資訊。", size=10)
    table(doc, ["候選池", "設定"], [
        ["檔數與流動性", f"每日約 {c['n_candidates']} 檔；20 日均成交值 ≥ {c['min_adv'] / 1e8:.0f} 億（8% 部位約 8000 萬，低於日成交值的 27%）"],
        ["分數組成", f"投信近 5 日淨買超占成交值 ×{c['w_trust5']}、近 10 日 ×{c['w_trust10']}；20 日相對強度 ×{c['w_rs20']}；"
                   f"成交值放大（5 日／20 日均量）×{c['w_valchg']}（皆為橫斷面 z 分數）"],
        ["差異化扣分", f"減去 {c['overlap_penalty']} × 出現在主動型 ETF 前 15 大的比例：越多主動型 ETF 重押的股票分數越低"],
        ["固定列入", "、".join(c["always_include"]) + "：重疊扣分會讓它永遠排不進前 30 名，但宣告允許 Agent 選它（上限 20%）"],
    ], widths=[85, 380])
    table(doc, ["入選依據", "資料來源（authority）", "例子"],
          [[b["label"], "、".join(b["authority"]), "；".join(b["examples"])] for b in decl["selection"]["allowed_bases"]],
          widths=[70, 90, 305])
    heading(doc, "Agent 的輸出與程式檢查", 2)
    bullets(doc, decl["selection"]["requirements"] + [
        "信心等級：" + "；".join(decl["selection"]["conviction_levels"].values()) + "。"])
    mapping_box(doc, [
        "observations：candidate_screen（候選分數與組成，排序用）、trust_flow（投信 5／10 日淨買超、連買天數）以及月營收、財報、法說會、重大訊息事件；"
        "topic 與依據 id 相同，來源 authority 如上表。",
        "inferences：每檔高信心部位的 premise_refs 指向該股的依據 observation，logic 說明「事實 → 為何高／中信心 → 與主題的關係」，"
        "counter_evidence 寫反方訊號；主動移出者另有一則移出推論。",
    ])


def _sec_discipline(doc, decl, S):
    r = decl["portfolio_rules"]
    heading(doc, "三、換股紀律（程式強制，違反就退回 Agent 重寫）", 2)
    bullets(doc, list(decl["discipline"][k] for k in ("keep_rank", "min_hold_days", "max_swaps", "stop_loss", "event_blackout")) + [
        "名單與信心等級沒有新證據就續抱；每次換股往返成本約 0.6%，候選分數變動不是換股理由。",
        f"狀態（名單、進場日、引用證據）存於 {decl['discipline']['state']}，隔日重建成 observation，續抱的推論可重新引用原證據。"])
    from .strategy.conviction import load_event_calendar
    evs = [e for e in load_event_calendar() if "2026-10-20" <= e["date"] <= "2026-11-30"]
    if evs:
        table(doc, ["初賽期間高波動事件（美國時間）", "前一交易日的處理"],
              [[f"{e['date']} {e['name']}" + ("（推估）" if e.get("estimated") else ""), "不新增、不升級，既有部位凍結"] for e in evs],
              widths=[230, 235])
    mapping_box(doc, [
        "observations：position_pnl（持有天數、相對成本報酬）與 event_calendar（事件排程）。",
        "inferences：停損出場以 position_pnl 為前提；事件前一交易日的 no_trade 推論引用 event_calendar。",
        "decisions：停損 → SELL_ALL；未滿 3 日者不會出現在賣出；事件前一交易日不會出現 BUY。",
    ])


def _p(x):
    return pct(x, 0, False)


def _sec_backtest(doc, decl, S):
    df, oos = S["default"], S["default_oos"]
    heading(doc, "四、代理回測摘要", 2)
    para(doc, f"LLM 無法回測，以「候選池分數前 11 名（前 5 名高、其餘中；已持有且名次 ≤ 30 者續抱）」代替 Agent 選股，其餘規則與實盤同一套程式。"
              f"自 {S['first']} 起每 {S['step']} 個交易日一個起點、每段 {S['length']} 個交易日（初賽約 24 日）、從 10 億現金建倉，"
              "計算相對「市值加權上限基準（2330 ≤ 20%、其他 ≤ 8%）」的超額報酬分布。", size=10)
    rows = []
    for lab, b in (("預設（遲滯續抱，高 8%／中 5.5%）", df), ("　其中樣本外", oos), ("敏感度：嚴格前 11 名（不續抱）", S.get("strict")),
                   ("敏感度：2330 固定高信心", S.get("tsmc_fixed")), ("對照：舊權重 9%／6%", S.get("old_weights")),
                   ("對照：合規基準 A（降級保底）", S.get("base_a"))):
        if not b:
            continue
        e = b["excess"]
        rows.append([lab, str(e["n"]), _p(e["p_gt_5"]), _p(e["p_gt_10"]), _p(e["p_lt_m10"]), pct(e["mean"]), pct(b["mdd"]["mean"]),
                     str(b.get("sum_drift_days", "—"))])
    table(doc, ["24 日窗口", "窗口數", "P(超額 > +5%)", "P(> +10%)", "P(< −10%)", "平均超額", "平均 MDD", "單檔 > 10%（天）"], rows,
          widths=[140, 40, 55, 50, 50, 45, 45, 45])
    bullets(doc, [
        f"預設版 24 日超額 5%～95% 區間為 {pct(df['excess']['p5'])} ～ {pct(df['excess']['p95'])}，平均換手 {df['turnover']['mean']:.1f} 倍、"
        f"費用與證交稅平均 {pct(df['cost']['mean'], 1, False)}：集中策略確實拉大與大盤的偏離，但左尾同樣變厚；"
        "不續抱、每天追逐候選名次的版本換手更高、結果更差，所以規則以「沒有新證據就續抱」為預設。",
        f"樣本外（窗口起點 ≥ {S['oos_start']}）只有 {oos['excess']['n']} 個彼此重疊的窗口，平均超額 {pct(oos['excess']['mean'])}；"
        "顯示策略的勝負高度依賴選股是否抓到當期主流。",
        f"違規檢查（逐日期末）：持股檔數 {df.get('sum_v_holdings', 0)}、現金 {df.get('sum_v_cash', 0)}、Active Share < 15% 共 {df.get('sum_ap_fail_days', 0)} 天；"
        f"單檔超過 10% 的 {df.get('sum_drift_days', 0)} 天中，買進當日即超過者 {df.get('sum_buy_over_cap_days', 0)} 天，其餘為持有中的價格漂移，隔日減回（官方給 5 個交易日調整）。"
        "高信心權重由 9% 下修為 8%，就是為了降低「買進當日收盤超過 10%」被認定為加碼違規的風險。",
        "名單偏誤：150 檔是 2026 年 7 月底依市值選出，之前的回測有前視與倖存者偏誤，動能／籌碼分數尤其容易買到事後的贏家；"
        f"主動型 ETF 持股用 {'、'.join(S.get('etf_data_dates', []))} 的真實前 15 大快照。回測不證明 Agent 能創造超額報酬，只能檢驗規則與成本。",
    ], size=9.5)
    mapping_box(doc, [
        "market_view.counter_evidence 與 inferences.counter_evidence：誠實寫出反方訊號；confidence 不得高於證據所能支持的程度。",
        "no_trade_decisions：沒有新證據或差距在無交易帶內的持股一律寫入，並引用對應推論。",
    ])


def _sec_fallback(doc, decl, S):
    fb = decl["fallback"]
    heading(doc, "五、降級保底", 2)
    para(doc, fb["when"], size=10)
    para(doc, f"{fb['name']}：{fb['definition']}", size=9.5)
    mapping_box(doc, [
        "market_view：降級日 regime／stance 為 neutral，logic 說明降級原因與方式（凍結或合規基準 A）。",
        "inferences／decisions：凍結時沿用原名單的推論並重新引用原證據；改用合規基準 A 時以「依合規基準 A 配置」為推論。",
    ])


def _sec_checklist(doc, decl, S):
    heading(doc, "六、每日 D-Plan 與本文件的對照檢查", 2)
    rows = [
        ["投資主題：集中持有 10–12 檔高信心股", "decisions.target_weight（高 8%／中 5.5%）；inferences 每檔一則", "conviction.build_portfolio"],
        ["投資理念：依公開事件選股", "observations.topic ∈ 五種依據；inferences 引用該 observation；decisions 引用 inference", "conviction_agent.check_picks"],
        ["每檔須有具名來源", "sources.authority 與 content_as_of；observations.source_ref；dplan_validate C1（引用鏈不可斷）", "validate()"],
        ["持有 3 日、每日換 4 檔、停損 12%、事件前一日", "position_pnl／event_calendar observation；SELL_ALL 與 BUY 的時點", "conviction.discipline"],
        ["Active Share ≥ 20%（官方 15%）", "target_weight 已通過 Active Share 檢查才進入 decisions", "conviction.build_portfolio"],
        ["現金與曝險承諾", "market_view.posture 由實際委託推得；dplan_validate C12", "plan.posture_for／validate()"],
        ["紀律：無證據不動", "no_trade_decisions 涵蓋所有未異動持股", "dplan_validate 覆蓋檢查"],
    ]
    table(doc, ["本文件的宣告", "每日 D-Plan 對應欄位", "程式檢查"], rows, widths=[120, 255, 90])
    para(doc, f"宣告版本 {decl['version']}（status = {decl['status']}）。{decl['status_note']}", size=8.5, italic=True)


def build_strategy_doc(out: Path | None = None, *, decl: dict | None = None, summary: dict | None = None,
                       template: Path | None = None, form_only: bool = False) -> Path:
    decl = decl or load_declaration()
    errs = validate_submission(decl)
    if errs:
        raise DocError("宣告檔不符官方格式限制：" + "；".join(errs))
    doc = docx.Document(str(template or TEMPLATE))
    _fill_form(doc, decl)
    if not form_only:
        S = summary or load_summary()
        _sec_structure(doc, decl, S)
        _sec_selection(doc, decl, S)
        _sec_discipline(doc, decl, S)
        _sec_backtest(doc, decl, S)
        _sec_fallback(doc, decl, S)
        _sec_checklist(doc, decl, S)
    p = out or OUTPUT_PATH
    p.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(p))
    return p
